import csv
from openpyxl import Workbook, load_workbook
from generators.locations import fix_location_name
import random
import os
import re
import io


def validate_text(value):
    if not value:
        return False

    return bool(re.fullmatch(r"[A-Za-z\s]+", value.strip()))

def parse_locations_file(file_obj, filename):
    file_obj.seek(0)
    name = (filename or "").lower()

    if name.endswith(".csv"):
        text = file_obj.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        reader.fieldnames = [(header or "").strip().lower() for header in reader.fieldnames]
        raw_rows = list(reader)
    elif name.endswith(".xlsx"):
        workbook = load_workbook(file_obj, data_only=True)
        sheet = workbook.active
        all_rows = list(sheet.iter_rows(values_only=True))
        headers = [(header or "").strip().lower() for header in all_rows[0]]
        raw_rows = []
        for values in all_rows[1:]:
            raw_rows.append(dict(zip(headers, values)))
    """ else:
        raise ValueError("Upload a .csv or .xlsx file.") """

    rows = []
    for row in raw_rows:
        zip_value = row["zip code"]
        if isinstance(zip_value, float) and zip_value.is_integer():
            zip_value = int(zip_value)
        zip_code = str(zip_value or "").strip()
        if zip_code:
            zip_code = zip_code.zfill(5)

        store_id = str(row["store id"] or "").strip()
        state = str(row["state"] or "").strip()
        county = str(row["county"] or "").strip()
        city = str(row["city"] or "").strip()

        if not store_id and not zip_code:
            continue

        rows.append({
            "store_id": store_id,
            "state": state,
            "county": county,
            "city": city,
            "zip_code": zip_code,
        })
    return rows

def load_subcategories(path):

    final_list = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)

        for row in reader:
            category = row["Category"]
            subcategory = row["Subcategory"]
            if category not in final_list:
                final_list[category] = [subcategory]
            else:
                final_list[category].append(subcategory)
    return final_list

def random_number_amount(): 
    rate = round(random.uniform(0, 0.1), 4)
    gross = round(random.uniform(50, 100000), 2)
    exempt = round(random.uniform(0, gross), 2)
    taxable = round((gross - exempt), 2)
    st_collected = round((taxable * rate), 2)

    percentage = round(random.uniform(0, 0.6), 4)
    tax_purchases = round((gross * percentage), 2)
    use_tax_accrued = round((tax_purchases * rate), 2)

    return gross, exempt, taxable, st_collected, tax_purchases, use_tax_accrued

def make_excel(subcategory_ids, num_transaction, pools, file_name=None):

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data import")

    headers = [
        "Store Id",
        "State",
        "County",
        "City",
        "Zip Code",
        "Subcategory",
        "Gross Sales",
        "Exempt sales",
        "Taxable sales",
        "Tax Collected",
        "Taxable purchases",
        "Use tax accrued",
    ]
    raw_list = []

    for t in range(num_transaction):
        gross, exempt, taxable, st_collected, tax_purchases, use_tax_accrued = random_number_amount()
        pool = pools[t % len(pools)]
        location = random.choice(pool["matches"])
        row = [
            pool["store_id"],
            location["state_name"],
            fix_location_name(location["county_name"]),
            fix_location_name(location["city"]),
            location["zip"],
            subcategory_ids[t % len(subcategory_ids)]["subcategory"],
            gross,
            exempt,
            taxable,
            st_collected,
            tax_purchases,
            use_tax_accrued,
        ]

        raw_list.append(row)


    ordered_list = sorted(raw_list, key=lambda x: (not x[0], str(x[0]).lower(), str(x[1]).lower(), str(x[2]).lower(), str(x[5]).lower()))
    ws.append(headers)
    for row in ordered_list:
        ws.append(row) 


    output_name = file_name

    os.makedirs("output", exist_ok = True)

    path = os.path.join("output", output_name)
    wb.save(path)

    return path

