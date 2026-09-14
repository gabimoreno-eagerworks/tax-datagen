import csv
import io
import random
import os
from generators.filing_calendar import make_short_name, STREET_NAMES
from openpyxl import Workbook, load_workbook
from generators.locations import fix_location_name
from generators.locations import fill_locations_for_rows, us_file, load_zip_rows

STATE_CODES = {
    "American Samoa": "AS",
    "District of Columbia": "DC",
    "Federated States of Micronesia": "FM",
    "Guam": "GU",
    "Marshall Islands": "MH",
    "Northern Mariana Islands": "MP",
    "Palau": "PW",
    "Puerto Rico": "PR",
    "Virgin Islands": "VI",
    "Alabama": "AL",
    "Alaska": "AK",
    "Arizona": "AZ",
    "Arkansas": "AR",
    "California": "CA",
    "Colorado": "CO",
    "Connecticut": "CT",
    "Delaware": "DE",
    "Florida": "FL",
    "Georgia": "GA",
    "Hawaii": "HI",
    "Idaho": "ID",
    "Illinois": "IL",
    "Indiana": "IN",
    "Iowa": "IA",
    "Kansas": "KS",
    "Kentucky": "KY",
    "Louisiana": "LA",
    "Maine": "ME",
    "Maryland": "MD",
    "Massachusetts": "MA",
    "Michigan": "MI",
    "Minnesota": "MN",
    "Mississippi": "MS",
    "Missouri": "MO",
    "Montana": "MT",
    "Nebraska": "NE",
    "Nevada": "NV",
    "New Hampshire": "NH",
    "New Jersey": "NJ",
    "New Mexico": "NM",
    "New York": "NY",
    "North Carolina": "NC",
    "North Dakota": "ND",
    "Ohio": "OH",
    "Oklahoma": "OK",
    "Oregon": "OR",
    "Pennsylvania": "PA",
    "Rhode Island": "RI",
    "South Carolina": "SC",
    "South Dakota": "SD",
    "Tennessee": "TN",
    "Texas": "TX",
    "Utah": "UT",
    "Vermont": "VT",
    "Virginia": "VA",
    "Washington": "WA",
    "West Virginia": "WV",
    "Wisconsin": "WI",
    "Wyoming": "WY",
}

def _decode_csv_bytes(raw):
    if isinstance(raw, str):
        return raw
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode the CSV file. Save it as UTF-8 and try again.")


def _csv_delimiter(header_line):
    semicolon_count = header_line.count(";")
    comma_count = header_line.count(",")
    tab_count = header_line.count("\t")
    if semicolon_count > comma_count and semicolon_count > tab_count:
        return ";"
    if tab_count > comma_count:
        return "\t"
    return ","


def parse_locations_file(file_obj, filename):
    file_obj.seek(0)
    name = (filename or "").lower()

    if name.endswith(".csv"):
        text = _decode_csv_bytes(file_obj.read())
        header_line = next((line for line in text.splitlines() if line.strip()), "")
        reader = csv.DictReader(io.StringIO(text), delimiter=_csv_delimiter(header_line))
        if not reader.fieldnames:
            raise ValueError("The CSV file has no header row.")
        reader.fieldnames = [(header or "").strip().lower() for header in reader.fieldnames]
        raw_rows = list(reader)
    elif name.endswith(".xlsx"):
        workbook = load_workbook(file_obj, data_only=True)
        sheet = workbook.active
        all_rows = list(sheet.iter_rows(values_only=True))
        headers = [(header or "").strip().lower() for header in all_rows[0]]
        raw_rows = []
        for values in all_rows[1:]:
            raw_rows.append(dict(zip(headers, values)))
    else:
        raise ValueError("Upload a .csv or .xlsx file.")

    if not raw_rows:
        raise ValueError("The file has no data rows.")

    headers_found = [header for header in raw_rows[0].keys() if header]
    if "entity name" not in headers_found or "state name" not in headers_found:
        raise ValueError(
            "Expected columns 'Entity Name' and 'State Name'. "
            f"Found: {headers_found}"
        )

    rows = []
    seen = set()
    for row in raw_rows:
        entity_name = str(row.get("entity name") or "").strip()
        state = str(row.get("state name") or "").strip()
        if not entity_name or not state:
            continue
        pair = (entity_name, state)
        if pair in seen:
            continue
        seen.add(pair)
        rows.append({
            "Entity Name": entity_name,
            "State": state,
        })
    return rows

def fill_outlets_for_rows(calendar_rows):
    zip_rows = load_zip_rows(us_file)
    outlets = []
    for row in calendar_rows:
        locations = fill_locations_for_rows(
            us_file, 3, state=row["State"], zip_rows=zip_rows
        )
        if not locations:
            continue
        for location in locations:
            outlets.append({
                "Entity Name": row["Entity Name"],
                "State": location["state"],
                "County": location["county"],
                "City": location["city"],
                "Zip Code": location["zip_code"],
            })
    return outlets

def get_state_code(outlet_state):
    return STATE_CODES.get(outlet_state, "XX")
  
def make_store_id(entity_name, outlet_state):
    entity_short_name = make_short_name(entity_name)
    outlet_state_code = get_state_code(outlet_state)
    return f"{random.randint(0, 99):02d}-{entity_short_name}{outlet_state_code}-{random.randint(0, 99):02d}"

def make_outlets_report(outlets_list, file_name=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Outlet Report"
    raw_list = []

    headers = [
        "Store ID",
        "Outlet Name",
        "State",
        "County",
        "City",
        "Zip Code",
        "Street Address",
        "Entity Name",
    ]
    ws.append(headers)

    for t in range(len(outlets_list)):
        outlet = outlets_list[t]
        store_id = make_store_id(outlet["Entity Name"], outlet["State"])
        number = random.randint(100, 9999)
        street_name = random.choice(STREET_NAMES)
        street_address = (
            f"{number} {street_name}, {outlet['City']}, "
            f"{outlet['State']} {outlet['Zip Code']}"
        )
        row = [
            store_id,
            store_id,
            outlet["State"],
            fix_location_name(outlet["County"]),
            fix_location_name(outlet["City"]),
            outlet["Zip Code"],
            street_address,
            outlet["Entity Name"],
        ]
        raw_list.append(row)

    ordered_list = sorted(raw_list, key=lambda x: (not x[0], str(x[0]).lower()))
    for row in ordered_list:
        ws.append(row) 

    output_name = file_name 
    os.makedirs("output", exist_ok=True)
    path = os.path.join("output", output_name)
    wb.save(path)

    return path